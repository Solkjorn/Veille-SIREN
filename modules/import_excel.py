import sqlite3
from contextlib import closing
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from config import BASE_SQLITE
from modules.societes_surveillees import initialiser_societes_surveillees
from modules.validation import siren_valide

NOM_FEUILLE = "Societes"
COLONNES_ATTENDUES = ("SIREN", "Actif", "Commentaire")


@dataclass(frozen=True)
class LigneImport:
    numero: int
    siren: str
    actif: bool
    commentaire: str


@dataclass(frozen=True)
class ErreurImport:
    numero: int
    message: str


@dataclass
class ApercuImport:
    lignes: list[LigneImport] = field(default_factory=list)
    erreurs: list[ErreurImport] = field(default_factory=list)


@dataclass(frozen=True)
class BilanImport:
    ajouts: int
    mises_a_jour: int
    lignes_ignorees: int
    erreurs: tuple[ErreurImport, ...]


def _normaliser_siren_excel(valeur) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, bool):
        return str(valeur)
    if isinstance(valeur, int):
        return str(valeur).zfill(9)
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur)).zfill(9)
    return str(valeur).strip().replace(" ", "")


def _normaliser_actif(valeur) -> bool:
    if isinstance(valeur, bool):
        return valeur
    if isinstance(valeur, (int, float)) and valeur in (0, 1):
        return bool(valeur)
    texte = str(valeur or "").strip().casefold()
    if texte in {"oui", "o", "true", "vrai", "1", "actif", "active"}:
        return True
    if texte in {"non", "n", "false", "faux", "0", "inactif", "inactive"}:
        return False
    raise ValueError("la valeur Actif doit indiquer oui ou non")


def _verifier_extension(nom_fichier: str) -> None:
    if Path(nom_fichier).suffix.casefold() != ".xlsx":
        raise ValueError("Le fichier doit être au format .xlsx")


def _analyser_classeur(classeur) -> ApercuImport:
    if NOM_FEUILLE not in classeur.sheetnames:
        raise ValueError(f"Feuille attendue absente : {NOM_FEUILLE}")
    feuille = classeur[NOM_FEUILLE]
    premiere_ligne = next(
        feuille.iter_rows(min_row=1, max_row=1, values_only=True), ()
    )
    entetes = tuple(str(valeur or "").strip() for valeur in premiere_ligne)
    while entetes and not entetes[-1]:
        entetes = entetes[:-1]
    if entetes != COLONNES_ATTENDUES:
        attendues = ", ".join(COLONNES_ATTENDUES)
        raise ValueError(f"Colonnes attendues : {attendues}")

    apercu = ApercuImport()
    sirens_vus = set()
    for numero, valeurs in enumerate(
        feuille.iter_rows(min_row=2, max_col=3, values_only=True), start=2
    ):
        if all(valeur is None for valeur in valeurs):
            continue
        siren = _normaliser_siren_excel(valeurs[0])
        erreurs_ligne = []
        if not siren_valide(siren):
            erreurs_ligne.append(f"SIREN invalide : {siren or '(vide)'}")
        elif siren in sirens_vus:
            erreurs_ligne.append(f"SIREN en double dans le fichier : {siren}")
        try:
            actif = _normaliser_actif(valeurs[1])
        except ValueError as erreur:
            erreurs_ligne.append(str(erreur))
            actif = False

        if erreurs_ligne:
            apercu.erreurs.extend(
                ErreurImport(numero, message) for message in erreurs_ligne
            )
            continue
        sirens_vus.add(siren)
        apercu.lignes.append(
            LigneImport(
                numero=numero,
                siren=siren,
                actif=actif,
                commentaire=str(valeurs[2] or "").strip(),
            )
        )
    return apercu


def analyser_fichier_excel(fichier_excel: Path) -> ApercuImport:
    """Analyse un classeur présent sur disque sans modifier la base."""
    chemin = Path(fichier_excel)
    _verifier_extension(chemin.name)
    if not chemin.is_file():
        raise FileNotFoundError(chemin)

    try:
        classeur = load_workbook(chemin, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as erreur:
        raise ValueError("Le fichier .xlsx est illisible ou corrompu") from erreur
    try:
        return _analyser_classeur(classeur)
    finally:
        classeur.close()


def analyser_televersement_excel(nom_fichier: str,
                                 contenu: bytes) -> ApercuImport:
    """Analyse en mémoire le contenu d'un fichier téléversé."""
    _verifier_extension(Path(nom_fichier).name)
    if not contenu:
        raise ValueError("Le fichier téléversé est vide")
    try:
        classeur = load_workbook(
            BytesIO(contenu), read_only=True, data_only=True
        )
    except (BadZipFile, InvalidFileException, OSError) as erreur:
        raise ValueError("Le fichier .xlsx est illisible ou corrompu") from erreur
    try:
        return _analyser_classeur(classeur)
    finally:
        classeur.close()


def importer_apercu(apercu: ApercuImport,
                    chemin: Path = BASE_SQLITE) -> BilanImport:
    """Importe atomiquement les lignes valides d'un aperçu."""
    initialiser_societes_surveillees(chemin)
    maintenant = datetime.now().isoformat(timespec="seconds")
    ajouts = 0
    mises_a_jour = 0

    with closing(sqlite3.connect(chemin)) as connexion:
        with connexion:
            for ligne in apercu.lignes:
                existe = connexion.execute(
                    "SELECT 1 FROM societes_surveillees WHERE siren = ?",
                    (ligne.siren,),
                ).fetchone()
                if existe:
                    connexion.execute(
                        """UPDATE societes_surveillees
                        SET actif = ?, commentaire = ?,
                            date_modification = ?, date_archivage = NULL
                        WHERE siren = ?""",
                        (int(ligne.actif), ligne.commentaire,
                         maintenant, ligne.siren),
                    )
                    mises_a_jour += 1
                else:
                    connexion.execute(
                        """INSERT INTO societes_surveillees
                        (siren, actif, commentaire, date_creation,
                         date_modification, date_archivage)
                        VALUES (?, ?, ?, ?, ?, NULL)""",
                        (ligne.siren, int(ligne.actif), ligne.commentaire,
                         maintenant, maintenant),
                    )
                    ajouts += 1

    return BilanImport(
        ajouts=ajouts,
        mises_a_jour=mises_a_jour,
        lignes_ignorees=len({erreur.numero for erreur in apercu.erreurs}),
        erreurs=tuple(apercu.erreurs),
    )


def importer_fichier_excel(fichier_excel: Path,
                           chemin: Path = BASE_SQLITE) -> BilanImport:
    return importer_apercu(analyser_fichier_excel(fichier_excel), chemin)
