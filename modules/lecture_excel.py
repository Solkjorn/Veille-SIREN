from openpyxl import load_workbook


def lire_sirens(fichier_excel):
    """
    Lit les SIREN du fichier Excel.
    Retourne une liste de dictionnaires.
    """

    wb = load_workbook(fichier_excel)
    ws = wb["Societes"]

    societes = []

    # On commence à la ligne 2 pour ignorer les en-têtes
    for ligne in ws.iter_rows(min_row=2, values_only=True):
        siren, actif, commentaire = ligne

        if siren is None:
            continue

        societes.append({
            "siren": str(siren),
            "actif": actif,
            "commentaire": commentaire
        })

    return societes
