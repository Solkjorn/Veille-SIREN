import tempfile
import unittest
from datetime import datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from modules.base_donnees import (
    alertes_immediates_actives, enregistrer_alertes, enregistrer_societe,
    lire_alertes,
)
from modules.comparaison import Changement
from modules.modele import Societe
from modules.societes_surveillees import (
    ajouter_societe_surveillee,
    supprimer_societe_surveillee,
)
from webapp import creer_application


class TestApplicationWeb(unittest.TestCase):
    def test_entetes_de_securite_locale(self):
        reponse = self.client.get("/")
        self.assertEqual(reponse.headers["X-Content-Type-Options"], "nosniff")

    def test_configuration_alertes_immediates_desactivee_par_defaut(self):
        self.assertFalse(alertes_immediates_actives(self.base))
        page = self.client.get("/alertes").get_data(as_text=True)
        self.assertIn("Envoyer immédiatement", page)
        reponse = self.client.post(
            "/alertes/configuration", data={"alertes_immediates": "1"}
        )
        self.assertEqual(reponse.status_code, 302)
        self.assertTrue(alertes_immediates_actives(self.base))
        self.assertEqual(reponse.headers["X-Frame-Options"], "DENY")
        self.assertEqual(reponse.headers["Referrer-Policy"], "no-referrer")
        self.assertIn("default-src 'self'", reponse.headers["Content-Security-Policy"])

    def test_affiche_filtre_et_traite_une_alerte(self):
        changement = Changement(
            champ="statut", libelle="Statut", ancienne_valeur="Active",
            nouvelle_valeur="Radiée", niveau="critique",
            categorie="cessation_radiation",
            regle="statut de cessation ou radiation",
        )
        enregistrer_alertes(
            "542051180", [changement], source="INSEE",
            date_detection=datetime(2026, 8, 3, 7, 5), chemin=self.base,
        )

        reponse = self.client.get("/alertes?niveau=critique&statut=nouvelle")
        page = reponse.get_data(as_text=True)
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("Alertes juridiques", page)
        self.assertIn("statut de cessation ou radiation", page)
        self.assertIn("Radiée", page)

        identifiant = lire_alertes(chemin=self.base)[0]["identifiant"]
        reponse = self.client.post(
            f"/alertes/{identifiant}/statut", data={"statut": "traitee"}
        )
        self.assertEqual(reponse.status_code, 302)
        self.assertIn(
            "Traitée",
            self.client.get("/alertes?statut=traitee").get_data(as_text=True),
        )

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name) / "web.sqlite"
        self.application = creer_application({
            "TESTING": True,
            "BASE_SQLITE": self.base,
            "CSRF_ACTIVE": False,
        })
        self.client = self.application.test_client()

    def tearDown(self):
        self.temp.cleanup()

    def test_tableau_de_bord_vide(self):
        reponse = self.client.get("/")
        self.assertEqual(reponse.status_code, 200)
        page = reponse.get_data(as_text=True)
        self.assertIn("Aucune société", page)
        self.assertIn('href="/imports"', page)
        self.assertNotIn("Sélectionner une liste de sociétés", page)

    def test_affiche_le_formulaire_d_import_sur_une_page_dediee(self):
        reponse = self.client.get("/imports")
        page = reponse.get_data(as_text=True)

        self.assertEqual(reponse.status_code, 200)
        self.assertIn("<h2>Imports</h2>", page)
        self.assertIn("Sélectionner une liste de sociétés", page)
        self.assertIn('accept=".xlsx"', page)
        self.assertIn("Télécharger le modèle", page)

    def test_telecharge_un_modele_excel_conforme(self):
        reponse = self.client.get("/imports/modele.xlsx")
        self.assertEqual(reponse.status_code, 200)
        self.assertIn(
            "modele_import_siren.xlsx", reponse.headers["Content-Disposition"]
        )
        classeur = load_workbook(BytesIO(reponse.data), read_only=True)
        try:
            self.assertEqual(classeur.sheetnames, ["Societes"])
            entetes = next(classeur["Societes"].iter_rows(values_only=True))
            self.assertEqual(entetes, ("SIREN", "Actif", "Commentaire"))
            self.assertEqual(classeur["Societes"].max_row, 1)
        finally:
            classeur.close()

    def test_tableau_de_bord_affiche_les_societes_et_indicateurs(self):
        ajouter_societe_surveillee(
            "542051180", commentaire="Active", chemin=self.base
        )
        enregistrer_societe(
            Societe(
                siren="542051180",
                raison_sociale="TOTALENERGIES SE",
                forme_juridique="Société européenne",
                capital="5 704 141 785 €",
                statut="Active",
                adresse="Courbevoie",
                dirigeant="Direction générale",
                derniere_publication_bodacc="10/07/2026",
                dernier_changement="Modification du capital",
                date_collecte=datetime(2026, 7, 30, 20, 0),
            ),
            self.base,
        )
        ajouter_societe_surveillee(
            "552032534", actif=False, commentaire="Inactive", chemin=self.base
        )
        ajouter_societe_surveillee("304154719", chemin=self.base)
        supprimer_societe_surveillee("304154719", chemin=self.base)

        reponse = self.client.get("/")
        page = reponse.get_data(as_text=True)

        self.assertEqual(reponse.status_code, 200)
        for texte in ("542051180", "552032534", "304154719",
                      "Active", "Inactive", "Archivée", "TOTALENERGIES SE",
                      "Modification du capital", "Mini-rapport"):
            self.assertIn(texte, page)
        self.assertIn("<th>Dernière modification</th>", page)
        self.assertIn("<th>Dernière collecte</th>", page)
        self.assertIn("10/07/2026", page)
        self.assertIn("30/07/2026 20:00", page)
        self.assertIn("Aucune collecte", page)
        self.assertIn('class="objet-modification"', page)
        self.assertIn('id="rapport-542051180"', page)
        self.assertIn('aria-expanded="false"', page)
        self.assertIn("Ouvrir la fiche société", page)
        self.assertIn("Fiche →", page)
        self.assertIn("/societes/542051180/historique", page)
        self.assertIn("Sociétés avec modifications", page)
        self.assertIn("Sociétés sans modification", page)

    def test_classe_une_societe_modifiee_dans_la_premiere_liste(self):
        ajouter_societe_surveillee("542051180", chemin=self.base)
        enregistrer_societe(
            Societe(
                siren="542051180", raison_sociale="TOTALENERGIES SE",
                adresse="Ancienne adresse", date_collecte=datetime(2026, 8, 1, 7),
            ), self.base,
        )
        enregistrer_societe(
            Societe(
                siren="542051180", raison_sociale="TOTALENERGIES SE",
                adresse="Nouvelle adresse", date_collecte=datetime(2026, 8, 2, 7),
            ), self.base,
        )

        page = self.client.get("/").get_data(as_text=True)
        debut_modifiees = page.index("Sociétés avec modifications")
        debut_stables = page.index("Sociétés sans modification")
        self.assertIn("TOTALENERGIES SE", page[debut_modifiees:debut_stables])
        self.assertNotIn("TOTALENERGIES SE", page[debut_stables:])

    def test_ajoute_modifie_et_change_l_etat_d_une_societe(self):
        reponse = self.client.post("/societes", data={
            "siren": "542051180",
            "actif": "1",
            "commentaire": "À suivre",
        })
        self.assertEqual(reponse.status_code, 302)

        self.client.post("/societes/542051180/modifier", data={
            "commentaire": "Commentaire modifié",
        })
        self.client.post("/societes/542051180/desactiver")
        page = self.client.get("/").get_data(as_text=True)
        self.assertIn("Commentaire modifié", page)
        self.assertIn("Inactive", page)

        self.client.post("/societes/542051180/activer")
        self.client.post("/societes/542051180/archiver")
        page = self.client.get("/").get_data(as_text=True)
        self.assertIn("Archivée", page)
        self.assertIn("Restaurer la société", page)

        self.client.post("/societes/542051180/restaurer")
        page = self.client.get("/").get_data(as_text=True)
        self.assertIn("Active", page)

    def test_protege_les_actions_avec_un_jeton_csrf(self):
        application = creer_application({
            "TESTING": True,
            "BASE_SQLITE": self.base,
            "SECRET_KEY": "test",
        })
        reponse = application.test_client().post("/societes", data={
            "siren": "542051180",
        })
        self.assertEqual(reponse.status_code, 400)

    def test_previsualise_puis_confirme_un_import_excel(self):
        classeur = Workbook()
        feuille = classeur.active
        feuille.title = "Societes"
        feuille.append(["SIREN", "Actif", "Commentaire"])
        feuille.append(["542051180", "Oui", "Importée"])
        feuille.append(["123", "Oui", "Invalide"])
        contenu = BytesIO()
        classeur.save(contenu)
        classeur.close()
        contenu.seek(0)

        reponse = self.client.post(
            "/import-excel/previsualiser",
            data={"fichier_excel": (contenu, "societes.xlsx")},
            content_type="multipart/form-data",
        )
        page = reponse.get_data(as_text=True)

        self.assertEqual(reponse.status_code, 200)
        self.assertIn("Prévisualisation", page)
        self.assertIn("542051180", page)
        self.assertIn("SIREN invalide", page)
        marqueur = 'name="identifiant_import"\n               value="'
        debut = page.index(marqueur) + len(marqueur)
        identifiant = page[debut:page.index('"', debut)]

        confirmation = self.client.post(
            "/import-excel/confirmer",
            data={"identifiant_import": identifiant},
            follow_redirects=True,
        )
        page_finale = confirmation.get_data(as_text=True)

        self.assertIn("Import terminé", page_finale)
        self.assertIn("<h2>Imports</h2>", page_finale)
        tableau_de_bord = self.client.get("/").get_data(as_text=True)
        self.assertIn("Importée", tableau_de_bord)

    def test_refuse_une_confirmation_d_import_expiree(self):
        reponse = self.client.post(
            "/import-excel/confirmer",
            data={"identifiant_import": "inconnu"},
            follow_redirects=True,
        )
        self.assertIn(
            "prévisualisation a expiré", reponse.get_data(as_text=True)
        )

    def test_affiche_l_historique_et_les_changements_d_une_societe(self):
        ajouter_societe_surveillee(
            "542051180", commentaire="Dossier prioritaire", chemin=self.base
        )
        enregistrer_societe(
            Societe(
                siren="542051180",
                raison_sociale="SOCIÉTÉ TEST",
                capital="100 000 €",
                statut="Active",
                date_collecte=datetime(2026, 7, 29, 10, 0),
            ),
            self.base,
        )
        enregistrer_societe(
            Societe(
                siren="542051180",
                raison_sociale="SOCIÉTÉ TEST",
                capital="120 000 €",
                statut="Active",
                source="INSEE + BODACC",
                derniere_publication_bodacc="Dépôt des comptes annuels 2025",
                provenance={"statut": "INSEE"},
                erreurs_sources={"INPI": "indisponible"},
                contradictions={"raison_sociale": {
                    "INSEE": "SOCIÉTÉ TEST", "BODACC": "Société Test",
                }},
                documents_inpi=[{
                    "identifiant": "acte-1", "type_document": "acte",
                    "date_depot": "2026-07-30", "libelle": "Statuts mis à jour",
                    "nom_document": "statuts", "confidentialite": "Public",
                }],
                date_collecte=datetime(2026, 7, 30, 10, 0),
            ),
            self.base,
        )
        enregistrer_alertes(
            "542051180",
            [Changement(
                "capital", "Capital", "100 000 €", "120 000 €",
                niveau="important", categorie="capital",
                regle="modification du capital",
            )],
            source="INPI", date_detection=datetime(2026, 7, 30, 10, 1),
            chemin=self.base,
        )

        reponse = self.client.get("/societes/542051180/historique")
        page = reponse.get_data(as_text=True)

        self.assertEqual(reponse.status_code, 200)
        self.assertIn("Historique des collectes", page)
        self.assertIn("29/07/2026 à 10:00", page)
        self.assertIn("30/07/2026 à 10:00", page)
        self.assertIn("Capital", page)
        self.assertIn("100 000 €", page)
        self.assertIn("120 000 €", page)
        self.assertIn("Sources et qualité des données", page)
        self.assertIn("INSEE + BODACC", page)
        self.assertIn("INPI : indisponible", page)
        self.assertIn("Contradictions détectées", page)
        self.assertIn("Fiche société", page)
        self.assertIn("Situation actuelle", page)
        self.assertIn("Vérifiée le 30/07/2026 à 10:00", page)
        self.assertIn("État administratif", page)
        self.assertIn("Registre national des entreprises (INPI)", page)
        self.assertIn("data.inpi.fr/entreprises/542051180", page)
        self.assertIn("Alerte important", page)
        self.assertIn("modification du capital", page)
        self.assertIn("30/07/2026 à 10:01", page)
        self.assertIn("Publication BODACC", page)
        self.assertIn("Dépôt des comptes annuels 2025", page)
        self.assertIn("Commentaire interne", page)
        self.assertIn("Dossier prioritaire", page)
        self.assertIn("Actes et comptes INPI", page)
        self.assertIn("Statuts mis à jour", page)
        self.assertIn("Télécharger le PDF", page)

        reponse = self.client.post(
            "/societes/542051180/preferences-alertes",
            data={"categories": ["dirigeant", "capital"]},
        )
        self.assertEqual(reponse.status_code, 302)
        page = self.client.get(
            "/societes/542051180/historique"
        ).get_data(as_text=True)
        self.assertIn("Préférences d'alertes", page)
        self.assertIn("value=\"capital\" checked", page)
        self.assertNotIn("value=\"publication\" checked", page)

    @patch(
        "webapp.SourceINPI.telecharger_document",
        return_value=(b"%PDF-1.4", "application/pdf"),
    )
    def test_telecharge_un_document_inpi_a_la_demande(self, mock_telecharger):
        ajouter_societe_surveillee("542051180", chemin=self.base)
        enregistrer_societe(Societe(
            "542051180", documents_inpi=[{
                "identifiant": "acte-1", "type_document": "acte",
                "nom_document": "statuts société", "libelle": "Statuts",
            }],
        ), self.base)
        reponse = self.client.get(
            "/documents-inpi/acte/acte-1/telecharger"
        )
        self.assertEqual(reponse.status_code, 200)
        self.assertEqual(reponse.mimetype, "application/pdf")
        self.assertIn("attachment", reponse.headers["Content-Disposition"])
        mock_telecharger.assert_called_once_with("acte", "acte-1")

    @patch(
        "webapp.SourceINPI.telecharger_document",
        return_value=(b'{"bilan":{"total":42}}', "application/json"),
    )
    def test_telecharge_les_donnees_de_compte_en_json(self, mock_telecharger):
        ajouter_societe_surveillee("542051180", chemin=self.base)
        enregistrer_societe(Societe(
            "542051180", documents_inpi=[{
                "identifiant": "bilan-1", "type_document": "bilan_saisi",
                "nom_document": "", "libelle": "S",
            }],
        ), self.base)

        reponse = self.client.get(
            "/documents-inpi/bilan_saisi/bilan-1/telecharger"
        )

        self.assertEqual(reponse.status_code, 200)
        self.assertEqual(reponse.mimetype, "application/json")
        self.assertIn(".json", reponse.headers["Content-Disposition"])
        self.assertIn(b'"total": 42', reponse.data)

    @patch(
        "webapp.SourceINPI.telecharger_document",
        return_value=(b'{"erreur":"document absent"}', "application/json"),
    )
    def test_ne_deguise_pas_une_reponse_json_en_pdf(self, mock_telecharger):
        ajouter_societe_surveillee("542051180", chemin=self.base)
        enregistrer_societe(Societe(
            "542051180", documents_inpi=[{
                "identifiant": "acte-1", "type_document": "acte",
                "nom_document": "acte", "libelle": "Acte",
            }],
        ), self.base)

        reponse = self.client.get(
            "/documents-inpi/acte/acte-1/telecharger"
        )

        self.assertEqual(reponse.status_code, 200)
        self.assertEqual(reponse.mimetype, "application/json")
        self.assertIn(".json", reponse.headers["Content-Disposition"])

    def test_historique_retourne_404_pour_un_siren_inconnu(self):
        self.assertEqual(
            self.client.get("/societes/542051180/historique").status_code,
            404,
        )

    def test_affiche_collectes_journal_et_rapports(self):
        ajouter_societe_surveillee("542051180", chemin=self.base)
        enregistrer_societe(
            Societe(
                siren="542051180",
                raison_sociale="SOCIÉTÉ TEST",
                statut="Active",
                source="Pappers",
            ),
            self.base,
        )

        page_collectes = self.client.get("/collectes").get_data(as_text=True)
        page_journal = self.client.get("/journal").get_data(as_text=True)
        page_rapports = self.client.get("/rapports").get_data(as_text=True)

        self.assertIn("Dernières collectes", page_collectes)
        self.assertIn("SOCIÉTÉ TEST", page_collectes)
        self.assertIn("Journal d'exécution", page_journal)
        self.assertIn("Rapports disponibles", page_rapports)
        self.assertIn("developpement_20260730.md", page_rapports)

    def test_consulte_et_telecharge_un_rapport_sans_traversee_de_chemin(self):
        consultation = self.client.get(
            "/rapports/developpement_20260730.md"
        )
        telechargement = self.client.get(
            "/rapports/developpement_20260730.md/telecharger"
        )

        self.assertEqual(consultation.status_code, 200)
        self.assertIn("Rapport journalier", consultation.get_data(as_text=True))
        self.assertEqual(telechargement.status_code, 200)
        self.assertIn("attachment", telechargement.headers["Content-Disposition"])
        telechargement.close()
        self.assertEqual(self.client.get("/rapports/inconnu.txt").status_code, 404)

    def test_consulte_un_rapport_html_dans_le_navigateur(self):
        dossier = Path(self.temp.name) / "rapports"
        dossier.mkdir()
        (dossier / "veille_test.html").write_text(
            "<!doctype html><title>Rapport HTML</title>", encoding="utf-8"
        )

        with patch("webapp.DOSSIER_RAPPORTS", dossier):
            liste = self.client.get("/rapports").get_data(as_text=True)
            consultation = self.client.get("/rapports/veille_test.html")

        self.assertIn("veille_test.html", liste)
        self.assertIn("Rapport HTML", liste)
        self.assertEqual(consultation.status_code, 200)
        self.assertEqual(consultation.mimetype, "text/html")
        self.assertIn("Rapport HTML", consultation.get_data(as_text=True))
        consultation.close()

    def test_erreurs_affiche_entreprises_diagnostic_et_rapport(self):
        dossier = Path(self.temp.name) / "rapports-erreurs"
        dossier.mkdir()
        rapport = dossier / "veille_erreurs.html"
        rapport.write_text(
            "<section class='erreurs'><ul><li><strong>542051180</strong> "
            "— Délai de lecture dépassé</li></ul></section>",
            encoding="utf-8",
        )
        ajouter_societe_surveillee("542051180", chemin=self.base)
        execution = {
            "id": 1, "type_execution": "hebdomadaire", "source": "pappers",
            "statut": "succes", "date_debut": "2026-08-02T07:00:00",
            "date_fin": "2026-08-02T07:10:00", "societes": 1,
            "modifications": 0, "erreurs": 1, "rapport": str(rapport),
            "message": "",
        }

        with patch("webapp.DOSSIER_RAPPORTS", dossier), patch(
            "webapp.lire_executions_veille", return_value=[execution]
        ), patch("webapp.lire_erreurs_taches_entre", return_value=[]):
            reponse = self.client.get("/erreurs")

        page = reponse.get_data(as_text=True)
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("542051180", page)
        self.assertIn("Délai de lecture dépassé", page)
        self.assertIn("Consulter le rapport complet", page)
        self.assertIn("/societes/542051180/historique", page)
        self.assertIn("/erreurs/542051180/reprendre", page)

    def test_formulaire_d_archivage_demande_confirmation(self):
        ajouter_societe_surveillee("542051180", chemin=self.base)
        page = self.client.get("/").get_data(as_text=True)
        self.assertIn('data-confirmation="Archiver cette société ?', page)

    @patch("webapp.lire_etat_tache")
    def test_affiche_la_planification(self, mock_etat):
        mock_etat.return_value = {
            "disponible": True,
            "etat": "Ready",
            "derniere": "",
            "prochaine": "2026-08-03T07:00:00",
            "resultat": 0,
        }
        page = self.client.get("/automatisation").get_data(as_text=True)
        self.assertIn("Automatisation hebdomadaire", page)
        self.assertIn("03/08/2026 à 07:00", page)
        self.assertIn("Lancer maintenant", page)

    @patch("webapp.lancer_tache")
    def test_lance_la_tache_depuis_interface(self, mock_lancer):
        reponse = self.client.post("/automatisation/lancer")
        self.assertEqual(reponse.status_code, 302)
        mock_lancer.assert_called_once_with()

    @patch("webapp.installer_tache")
    def test_modifie_la_planification_depuis_interface(self, mock_installer):
        reponse = self.client.post("/automatisation/configurer", data={
            "jour": "2", "heure": "9", "minute": "30", "source": "inpi",
        })
        self.assertEqual(reponse.status_code, 302)
        mock_installer.assert_called_once_with(
            source="inpi", jour_semaine=2, heure=9, minute=30
        )

    def test_pages_de_consolidation_sont_accessibles(self):
        for route, titre in (
            ("/documents", "Documents juridiques"),
            ("/recherche", "Recherche globale"),
            ("/audit", "Journal d'audit"),
            ("/diagnostic", "Diagnostic de l'application"),
        ):
            reponse = self.client.get(route)
            self.assertEqual(reponse.status_code, 200, route)
            self.assertIn(titre, reponse.get_data(as_text=True))

    def test_configure_le_destinataire_d_un_portefeuille(self):
        from modules.base_donnees import creer_portefeuille, lire_portefeuilles
        identifiant = creer_portefeuille("Clients", chemin=self.base)
        reponse = self.client.post(
            f"/portefeuilles/{identifiant}/destinataire",
            data={"destinataire": "clients@test.fr"},
        )
        self.assertEqual(reponse.status_code, 302)
        self.assertEqual(
            lire_portefeuilles(self.base)[0]["destinataire"], "clients@test.fr"
        )


if __name__ == "__main__":
    unittest.main()
